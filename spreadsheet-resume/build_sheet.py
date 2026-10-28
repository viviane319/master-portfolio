"""Builds viv-tran-cv.xlsx: a tabbed, spreadsheet-style CV.

Edit the data below and run:  python build_sheet.py
(needs openpyxl:  pip install openpyxl)
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAME = "Viv Tran"

ABOUT = [
    ("Name", "Viv Tran (she/her)"),
    ("Based in", "Brooklyn, NY"),
    ("Now", "Marketing & Communications Manager, eimmigration · Founder, Do Better Comms"),
    ("In one line", "Narrative strategist and community mobilizer helping mission-driven organizations tell their stories, build community, and make an impact."),
    ("Bio", "For the past decade, Viv has worked at the intersection of philanthropy, policy, and tech to advance social good. "
            "At eimmigration she connects immigration lawyers with tools that help them show up for their communities. "
            "Before that she co-directed communications at Donors of Color Network, where her campaign work helped move $120M to BIPOC-led climate groups in two years. "
            "A National Urban Fellow and Georgetown Master of Policy Management, Viv is also trained in Eastern massage, meditation, and altar building, "
            "practices that inform her holistic, global approach to the work."),
    ("Approach", "Put people at the center. Use storytelling and community building to make industry more human, more accessible, and more connected."),
    ("LinkedIn", "https://www.linkedin.com/in/viv-tran/"),
]

# (org, role, years, location, [highlights])
WORK = [
    ("eimmigration", "Marketing & Communications Manager", "2025 –", "New York, NY", [
        "Brand narratives and digital campaigns that translate immigration lawyers' needs and goals into audience-centric messaging",
        "Customer advocacy: referral program, customer advisory board (CAB), customer events, and listening initiatives",
        "Regional and national field marketing: sponsorships, industry events, customer experiences",
        "Multi-platform review program activating customer evangelists across the immigration law community",
    ]),
    ("Donors of Color Network / Action", "Communications Co-Director, Brand & Operations", "2022 – 2025", "New York, NY / Austin, TX", [
        "Led external and internal multi-channel communications: campaigns, funder partnerships, coalition building, donor engagement",
        "Brand strategy for four separate brands: visual identities, storytelling, key messaging",
        "Thought leadership and press: organization and executive placements in Forbes, Candid, Nonprofit Quarterly, Philanthropy News Digest",
        "Climate Funders Justice Pledge (CFJP) campaign communications, influencing top climate funders: $120M moved to BIPOC-led climate groups in two years",
        "Data-driven campaigns with TargetSmart and Blackbaud",
        "Managed budgets, timelines, staff, vendors, and consultants across all communications projects",
    ]),
    ("Grenzebach Glier & Associates (GG+A)", "Senior Program Manager, Philanthropic Analytics Practice + DEI", "2020 – 2022", "Chicago, IL (remote)", [
        "Consulted 25+ nonprofits on campaign performance, messaging, and fundraising through data-driven research and surveys",
        "Thought leadership, marketing campaigns, and integrated advertising for the GG+A SurveyLab research tool",
        "Industry-wide survey and data report on COVID-19 fundraising impacts and global trends",
        "Co-developed firm-wide DEI strategy: communications plan, inclusive hiring practices, pulse survey, bias training; served on diversity council",
        "Drafted 20+ client proposals and contracts with business development and finance",
    ]),
    ("Energy Policy Institute at UChicago (EPIC)", "Communications Manager (Contract)", "2020", "Chicago, IL", [
        "Led external communications and integrated media campaigns translating research into multi-channel content",
        "Built targeted press lists for research releases; tracked media coverage and paid advertising performance",
        "Wrote and edited semi-weekly and monthly newsletters on media, partner, and research news",
    ]),
    ("SmithBucklin", "Nonprofit Marketing Operations Associate", "2018 – 2020", "Chicago, IL", [
        "Program-managed two annual silent auctions raising $110K+ for local charities in Orlando and Austin",
        "Owned the advertising program: $72K in paid and sponsored ads",
        "Managed programming for four conferences; $400K in event budget forecasting and reconciliation",
        "Facilitated a volunteer special interest group, two mentorship programs, and eight councils of 120+ people",
        "Counseled two boards of directors on board meetings and event partnerships",
    ]),
    ("The Light House DaNang Hostel", "Co-Founder: Management & Social Impact", "2017 –", "Da Nang, Vietnam", [
        "Lead-generation pipeline: 200+ bookings in the first month, sold out in the first week",
        "Secured $10K in investment capital through partnerships with local community leaders",
        "Negotiated real-estate contracts; managed design and operating budget",
    ]),
    ("Oracle", "Cloud Applications Consultant", "2016 – 2017", "Austin, TX / Redwood City, CA", [
        "Built cloud implementation road maps for C-level executives at SMB companies",
        "Account strategy with virtual sales teams across all CX applications (FY17 attainment: Q1 88%, Q2 104%, Q3 92%)",
        "Graduate of the Oracle Sales Academy",
    ]),
    ("Cardinal Health", "Communications & Events Intern", "2016", "Dublin, OH / Chicago, IL", [
        "Supported registration, logistics, and live updates for an 8,000-attendee conference and trade show",
        "Organized 300+ show-floor booths; wrote the internal publication for the 2016 intern class",
    ]),
    ("Ardina", "Community & Social Media Manager", "2016", "Columbus, OH", [
        "Ran the digital community: social media, email marketing, monthly campaigns tracked in Google Analytics and HubSpot",
    ]),
    ("614 Media Group", "Marketing Intern", "2015", "Columbus, OH", [
        "Partner coordination for Columbus Restaurant Week and ColumBEST; content and photography for 614 and Uweekly",
    ]),
    ("MediaSource", "Media Measurement & Analysis Intern", "2015", "Columbus, OH", [
        "Seasonal media-coverage report for the American Dairy Association; campaign performance vs. client goals",
    ]),
    ("Ohio State School of Communication", "News Research Assistant Analyst", "2015", "Columbus, OH", [
        "Content analysis of 100+ television newscasts using an original codebook",
    ]),
    ("Studio 614", "Art Intern", "2015", "Columbus, OH", [
        "Taught painting classes for community events and fundraisers; made promotional video, blog, and social content",
    ]),
    ("Hola Columbus / Hola Media Group", "Social Media Manager", "2013 – 2015", "Columbus, OH", [
        "Managed social media and website for Hola Columbus and its advertisers; on-set live posting for TV",
    ]),
    ("APIA Vote Michigan", "Political Advocacy Intern", "2012", "Detroit, MI", [
        "Voter-engagement programs, voter drives, and workshops on leadership and civic participation",
    ]),
    ("88.1 WSDP", "DJ & Host", "2011 – 2012", "Plymouth, MI", [
        "Hosted an afternoon show; interviews for the Community Focus segment; news writing for The News File",
    ]),
]

CONSULTING = [
    ("Do Better Comms", "Founder: creative communications agency", "2025 –", "New York, NY", [
        "Partners with mission-driven organizations to make messaging more accessible, memorable, and undeniable",
        "Messaging and mission-statement rewrites: no jargon, no TL;DR",
        "Brand narrative, campaign strategy, and storytelling for nonprofits, advocacy groups, schools, and libraries",
    ]),
]

# (org, role, years, location, note)
FELLOWSHIPS = [
    ("National Urban Fellows", "Graduate Policy Management Fellow (Class of 2025)", "2024 – 2025", "Washington, D.C. / New York, NY", "Fully funded M.P.M. at Georgetown McCourt; residency, executive mentorship, and leadership coaching"),
    ("American YouthWorks", "Mentor & GED Tutor", "2022 – 2024", "Austin, TX", ""),
    ("Chicago Women in Philanthropy", "Emerging Leader", "2020 – 2021", "Chicago, IL", ""),
    ("UN Major Group for Children and Youth", "Campaign Ambassador", "2020 – 2021", "global", ""),
    ("APIA Vote", "Civic Engagement Volunteer", "2012 –", "national", "Volunteer since the 2012 internship"),
    ("Columbus Creative Industry Mixer", "Board of Directors", "2016", "Columbus, OH", ""),
    ("TEDxOhioStateUniversity", "Communications & Planning Committee", "2015 – 2016", "Columbus, OH", "Co-designed the 2016 media kit and pamphlet; liaison to performers and guests"),
    ("GiveBackHack", "Project Hygeia: CIO", "2016", "Columbus, OH", "Subscription site and hotel partnership for a soap-recycling venture"),
    ("Alleviating Poverty Through Entrepreneurship Summit", "Event Organizer", "2014 – 2016", "Columbus, OH", ""),
]

# (year, place, title, role, link)
EVENTS = [
    ("2026", "New York, NY", "Dear Crisis, Sincerely… Grief: climate grief altar-making and reflection, NY Climate Week (Caravan of Dreams)", "Co-host with Kayla Palisoc", "https://dearcrisis.com/"),
    ("2026", "online", "Do Better Comms launch", "Founder", ""),
    ("2020", "global", "GG+A report: COVID-19 impacts on fundraising", "Co-author / producer", ""),
]

# (outlet, year, note)
PRESS = [
    ("Forbes", "2022 – 2025", "Donors of Color Network organization and executive placements"),
    ("Candid", "2022 – 2025", "Donors of Color Network placements"),
    ("Nonprofit Quarterly", "2022 – 2025", "Donors of Color Network placements"),
    ("Philanthropy News Digest", "2022 – 2025", "Donors of Color Network placements"),
]

IMPACT = [
    ("$120M", "moved to BIPOC-led climate groups in two years (CFJP campaign, DOCN)"),
    ("25+", "nonprofits advised on messaging, campaigns, and fundraising (GG+A)"),
    ("4", "brands built or refreshed: identities, storytelling, messaging (DOCN)"),
    ("$400K", "in event budgets forecast and reconciled (SmithBucklin)"),
    ("$110K+", "raised through silent auctions for local charities (SmithBucklin)"),
    ("200+", "bookings in a hostel's first month, sold out week one (The Light House)"),
]

SKILLS = [
    ("Strategy", "Brand & narrative strategy · Campaign strategy · Messaging · Thought leadership · Press & media relations · Coalition building"),
    ("Community", "Customer & donor engagement · Advocacy programs · Events & convenings · Facilitation · Field marketing"),
    ("Research & data", "Survey design · Data-driven campaigns · Content analysis · Media measurement"),
    ("CRM & fundraising tools", "HubSpot · Salesforce · Raiser's Edge · EveryAction · Benevity · Qualtrics · Oracle CX · Zendesk · GiveSmart · Cvent"),
    ("Digital & design", "Adobe Suite · Canva · WordPress · Squarespace · Google Analytics · Mailchimp · Hootsuite"),
    ("Work tools", "Microsoft Suite · Google Workspace · Asana · ClickUp · Trello · Teamwork"),
    ("Practices", "Eastern massage · Meditation · Altar building"),
]

EDUCATION = [
    ("Georgetown University, McCourt School of Public Policy", "Master of Policy Management (M.P.M.)", "2025", "Washington, D.C.", "Studied the intersection of philanthropy and global policy; National Urban Fellow"),
    ("The Ohio State University", "B.A., Strategic Communications", "", "Columbus, OH", ""),
]

# ---------------------------------------------------------------- styling
ARIAL = "Arial"
# Altarism palette + Tulsi Green
EMBER, CREAM_HEX, OXBLOOD, AMBER, PLUM, ASH, TULSI = "180F06", "F6F3EC", "6E413A", "B77C45", "755367", "AAA7A1", "566B47"
BLACK = PatternFill("solid", fgColor=EMBER)
CREAM = PatternFill("solid", fgColor=CREAM_HEX)
OX_FILL = PatternFill("solid", fgColor=OXBLOOD)
TAB_COLORS = [OXBLOOD, PLUM, AMBER, TULSI, ASH, EMBER, OXBLOOD, PLUM, AMBER]
THIN = Side(style="thin", color=EMBER)
TOP = Border(top=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTHS = {"A": 34, "B": 38, "C": 13, "D": 22, "E": 70, "F": 30}


def sheet(wb, title, heading):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = TAB_COLORS[len(wb.sheetnames) - 1]
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws["A1"] = heading
    ws["A1"].font = Font(name=ARIAL, bold=True, size=11, color=OXBLOOD)
    ws.merge_cells("A1:F1")
    ws["A1"].border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.sheet_view.showGridLines = True
    ws.freeze_panes = "A3"
    return ws


def band(ws, row, text):
    ws.cell(row, 1, text)
    for c in range(1, 7):
        ws.cell(row, c).fill = BLACK
        ws.cell(row, c).font = Font(name=ARIAL, bold=True, color=CREAM_HEX)
    return row + 1


def put(ws, row, col, value, bold=False, italic=False, fill=None, link=False):
    cell = ws.cell(row, col, value)
    cell.font = Font(name=ARIAL, bold=bold, italic=italic, underline="single" if link else None, color=PLUM if link else EMBER)
    cell.alignment = WRAP
    if link:
        cell.hyperlink = value
    if fill:
        cell.fill = fill
    return cell


def roles(ws, row, items):
    for org, role, years, place, bullets in items:
        first = row
        put(ws, row, 1, org, bold=True)
        put(ws, row, 2, role)
        put(ws, row, 3, years)
        put(ws, row, 4, place)
        for b in bullets:
            put(ws, row, 5, b)
            row += 1
        if row - first > 1:
            for col in (1, 2, 3, 4):
                ws.merge_cells(start_row=first, start_column=col, end_row=row - 1, end_column=col)
        for c in range(1, 7):
            ws.cell(first, c).border = TOP
    return row


wb = Workbook()
intro = wb.active
intro.title = "Intro"
intro.column_dimensions["A"].width = 110
intro.row_dimensions[1].height = 360
intro["A1"] = f"To view {NAME}'s CV, please click the tabs below:"
intro["A1"].font = Font(name=ARIAL, size=20, color=CREAM_HEX)
intro["A1"].fill = OX_FILL
intro.sheet_properties.tabColor = TAB_COLORS[0]
intro["A1"].alignment = Alignment(horizontal="center", vertical="center")
intro.sheet_view.showGridLines = False

ws = sheet(wb, "About", "About")
r = band(ws, 2, "About")
for label, text in ABOUT:
    put(ws, r, 1, label, bold=True)
    put(ws, r, 2, text, link=text.startswith("http"), fill=CREAM if label == "In one line" else None)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    if label == "Bio":
        ws.row_dimensions[r].height = 95
    r += 1

ws = sheet(wb, "Work", "Work Experience")
r = band(ws, 2, "Work")
put(ws, r, 2, "(for more about this work, please see the microsite and experience board in this repo)", italic=True, fill=CREAM)
put(ws, r, 5, "For impact numbers, please check out the \"Impact & Press\" tab below", italic=True, fill=CREAM)
r = roles(ws, r + 1, WORK)

ws = sheet(wb, "Consulting", "Consulting")
r = band(ws, 2, "Consulting")
r = roles(ws, r, CONSULTING)

ws = sheet(wb, "Fellowships, Leadership", "Fellowships, Leadership & Service")
r = band(ws, 2, "Fellowships, Leadership & Service")
for org, role, years, place, note in FELLOWSHIPS:
    put(ws, r, 1, org, bold=True)
    put(ws, r, 2, role)
    put(ws, r, 3, years)
    put(ws, r, 4, place)
    put(ws, r, 5, note)
    for c in range(1, 7):
        ws.cell(r, c).border = TOP
    r += 1

ws = sheet(wb, "Events & Talks", "Events, Talks & Publications")
r = band(ws, 2, "Events, Talks & Publications")
for year, place, title, role, link in EVENTS:
    put(ws, r, 1, role, bold=True)
    put(ws, r, 3, year)
    put(ws, r, 4, place)
    put(ws, r, 5, title, fill=CREAM)
    if link:
        put(ws, r, 6, link, link=True)
    r += 1

ws = sheet(wb, "Impact & Press", "Impact & Press")
r = band(ws, 2, "Impact")
for num, text in IMPACT:
    put(ws, r, 1, num, bold=True)
    put(ws, r, 2, text)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r += 1
r = band(ws, r + 1, "Press")
for outlet, years, note in PRESS:
    put(ws, r, 1, outlet, bold=True)
    put(ws, r, 3, years)
    put(ws, r, 5, note)
    r += 1

ws = sheet(wb, "Skills & Tools", "Skills & Tools")
r = band(ws, 2, "Skills & Tools")
for label, text in SKILLS:
    put(ws, r, 1, label, bold=True)
    put(ws, r, 2, text)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 1

ws = sheet(wb, "Education", "Education")
r = band(ws, 2, "Education")
for school, degree, year, place, note in EDUCATION:
    put(ws, r, 1, school, bold=True)
    put(ws, r, 2, degree)
    put(ws, r, 3, year)
    put(ws, r, 4, place)
    put(ws, r, 5, note)
    r += 1

wb.save("viv-tran-cv.xlsx")
print("wrote viv-tran-cv.xlsx")
